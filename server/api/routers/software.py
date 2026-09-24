"""
Software management: uninstall (agent-side) + deployment (upload/URL → push to agents).
Also provides software inventory listing and XLSX export.
"""
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Query
from fastapi.responses import JSONResponse, StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, text
from typing import Optional, List
import uuid, os, hashlib, json as _json, io
from datetime import datetime, timezone

from api.database import get_db
from api.models.models import Agent
from api.services.auth import get_current_user, get_agent_by_api_key
from fastapi import Header

async def get_current_agent(
    x_api_key: str = Header(..., alias="X-API-Key"),
    db: AsyncSession = Depends(get_db),
) -> Agent:
    from fastapi import HTTPException
    agent = await get_agent_by_api_key(db, x_api_key)
    if not agent:
        raise HTTPException(status_code=401, detail="Invalid API key")
    return agent

router = APIRouter()

DEPLOY_PACKAGES_DIR = "/app/static/deploy-packages"
os.makedirs(DEPLOY_PACKAGES_DIR, exist_ok=True)


# ─── DB helpers (raw SQL, no new ORM models needed yet) ──────────────────────

async def _ensure_tables(db: AsyncSession):
    """Idempotent table creation for software deploy tables."""
    await db.execute(text("""
        CREATE TABLE IF NOT EXISTS software_uninstall_jobs (
            id                   UUID PRIMARY KEY DEFAULT gen_random_uuid(),
            agent_id             UUID REFERENCES agents(id) ON DELETE CASCADE,
            software_name        TEXT NOT NULL,
            software_version     TEXT DEFAULT '',
            status               TEXT DEFAULT 'pending',
            output               TEXT DEFAULT '',
            error_message        TEXT,
            triggered_by         UUID,
            triggered_by_username TEXT,
            queued_at            TIMESTAMPTZ DEFAULT NOW(),
            finished_at          TIMESTAMPTZ
        )
    """))
    await db.execute(text("""
        CREATE TABLE IF NOT EXISTS software_packages (
            id UUID PRIMARY KEY DEFAULT gen_random_uuid(),
            name TEXT NOT NULL,
            version TEXT,
            description TEXT,
            file_path TEXT,
            download_url TEXT,
            checksum_sha256 TEXT,
            installer_type TEXT DEFAULT 'exe',
            install_args TEXT DEFAULT '',
            os_type TEXT DEFAULT 'windows',
            size_bytes BIGINT,
            created_by UUID,
            created_at TIMESTAMPTZ DEFAULT NOW()
        )
    """))
    await db.execute(text("""
        CREATE TABLE IF NOT EXISTS software_deploy_jobs (
            id UUID PRIMARY KEY DEFAULT gen_random_uuid(),
            package_id UUID,
            agent_id UUID REFERENCES agents(id) ON DELETE CASCADE,
            status TEXT DEFAULT 'pending',
            output TEXT DEFAULT '',
            error_message TEXT,
            triggered_by UUID,
            triggered_by_username TEXT,
            queued_at TIMESTAMPTZ DEFAULT NOW(),
            started_at TIMESTAMPTZ,
            finished_at TIMESTAMPTZ
        )
    """))
    await db.commit()


# ─── Package catalog ──────────────────────────────────────────────────────────

@router.get("/software-deploy/packages")
async def list_packages(db: AsyncSession = Depends(get_db), _=Depends(get_current_user)):
    await _ensure_tables(db)
    rows = await db.execute(text(
        "SELECT id, name, version, description, download_url, file_path, checksum_sha256, "
        "installer_type, install_args, os_type, size_bytes, created_at FROM software_packages ORDER BY created_at DESC"
    ))
    result = []
    for r in rows.fetchall():
        result.append({
            "id": str(r[0]), "name": r[1], "version": r[2], "description": r[3],
            "download_url": r[4], "file_path": r[5], "checksum_sha256": r[6],
            "installer_type": r[7], "install_args": r[8], "os_type": r[9],
            "size_bytes": r[10], "created_at": r[11].isoformat() if r[11] else None,
        })
    return result


@router.post("/software-deploy/packages")
async def create_package(body: dict, db: AsyncSession = Depends(get_db), user=Depends(get_current_user)):
    """Register a URL-based package (no file upload)."""
    await _ensure_tables(db)
    name = body.get("name", "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="name is required")
    download_url = body.get("download_url", "").strip()
    if not download_url:
        raise HTTPException(status_code=400, detail="download_url is required for URL-based packages")

    pkg_id = str(uuid.uuid4())
    await db.execute(text("""
        INSERT INTO software_packages
            (id, name, version, description, download_url, checksum_sha256,
             installer_type, install_args, os_type, created_by)
        VALUES (:id, :name, :version, :desc, :url, :checksum,
                :type, :args, :os, CAST(:uid AS uuid))
    """), {
        "id": pkg_id, "name": name, "version": body.get("version"),
        "desc": body.get("description"), "url": download_url,
        "checksum": body.get("checksum_sha256"), "type": body.get("installer_type", "exe"),
        "args": body.get("install_args", ""), "os": body.get("os_type", "windows"),
        "uid": str(user.id),
    })
    await db.commit()
    return {"id": pkg_id, "status": "created"}


@router.post("/software-deploy/packages/upload")
async def upload_package(
    name: str = Form(...),
    version: str = Form(None),
    description: str = Form(None),
    installer_type: str = Form("exe"),
    install_args: str = Form(""),
    os_type: str = Form("windows"),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    user=Depends(get_current_user),
):
    await _ensure_tables(db)
    pkg_id = str(uuid.uuid4())
    pkg_dir = os.path.join(DEPLOY_PACKAGES_DIR, pkg_id)
    os.makedirs(pkg_dir, exist_ok=True)

    dest_path = os.path.join(pkg_dir, file.filename)
    content = await file.read()
    with open(dest_path, "wb") as f:
        f.write(content)

    checksum = hashlib.sha256(content).hexdigest()
    size = len(content)

    await db.execute(text("""
        INSERT INTO software_packages
            (id, name, version, description, file_path, checksum_sha256,
             installer_type, install_args, os_type, size_bytes, created_by)
        VALUES (:id, :name, :version, :desc, :path, :checksum,
                :type, :args, :os, :size, CAST(:uid AS uuid))
    """), {
        "id": pkg_id, "name": name, "version": version, "desc": description,
        "path": dest_path, "checksum": checksum, "type": installer_type,
        "args": install_args, "os": os_type, "size": size, "uid": str(user.id),
    })
    await db.commit()
    return {"id": pkg_id, "checksum_sha256": checksum, "size_bytes": size, "status": "uploaded"}


@router.delete("/software-deploy/packages/{package_id}")
async def delete_package(package_id: str, db: AsyncSession = Depends(get_db), _=Depends(get_current_user)):
    row = await db.execute(text("SELECT file_path FROM software_packages WHERE id = CAST(:id AS uuid)"), {"id": package_id})
    pkg = row.fetchone()
    if not pkg:
        raise HTTPException(status_code=404, detail="Package not found")
    # Remove file if uploaded
    if pkg[0] and os.path.exists(pkg[0]):
        try:
            import shutil
            shutil.rmtree(os.path.dirname(pkg[0]), ignore_errors=True)
        except Exception:
            pass
    await db.execute(text("DELETE FROM software_packages WHERE id = CAST(:id AS uuid)"), {"id": package_id})
    await db.commit()
    return {"status": "deleted"}


# ─── Deployment ───────────────────────────────────────────────────────────────

@router.post("/software-deploy/deploy")
async def deploy_package(body: dict, db: AsyncSession = Depends(get_db), user=Depends(get_current_user)):
    """Deploy a package to one or more agents."""
    await _ensure_tables(db)
    package_id = body.get("package_id")
    agent_ids = body.get("agent_ids", [])
    if not package_id or not agent_ids:
        raise HTTPException(status_code=400, detail="package_id and agent_ids required")

    # Load package
    row = await db.execute(text(
        "SELECT name, version, download_url, file_path, checksum_sha256, installer_type, install_args "
        "FROM software_packages WHERE id = CAST(:id AS uuid)"
    ), {"id": package_id})
    pkg = row.fetchone()
    if not pkg:
        raise HTTPException(status_code=404, detail="Package not found")

    name, version, download_url, file_path, checksum, installer_type, install_args = pkg

    # Resolve URL: if file_path (uploaded), build internal URL
    server_url_row = await db.execute(text("SELECT value FROM system_settings WHERE key = 'general'"))
    server_url_setting = server_url_row.fetchone()
    server_base = "http://kifaa.kenyanut.com"
    if server_url_setting:
        try:
            import json
            settings_data = json.loads(server_url_setting[0])
            # Try to get server URL from settings if stored
        except Exception:
            pass

    url = download_url
    if not url and file_path:
        # Derive URL from file_path: /app/static/deploy-packages/{id}/{filename}
        relative = file_path.replace("/app/static/", "")
        url = f"{server_base}/{relative}"

    if not url:
        raise HTTPException(status_code=400, detail="Package has no download URL or file")

    job_ids = []
    for agent_id in agent_ids:
        job_id = str(uuid.uuid4())
        await db.execute(text("""
            INSERT INTO software_deploy_jobs
                (id, package_id, agent_id, status, triggered_by, triggered_by_username)
            VALUES (:id, CAST(:pkg AS uuid), CAST(:aid AS uuid), 'pending',
                    CAST(:uid AS uuid), :uname)
        """), {
            "id": job_id, "pkg": package_id, "aid": agent_id,
            "uid": str(user.id), "uname": user.username,
        })
        payload = _json.dumps({
            "deploy_job_id": job_id,
            "name": name,
            "version": version,
            "url": url,
            "args": install_args or "",
            "checksum": checksum or "",
            "installer_type": installer_type or "exe",
        })
        await db.execute(text("""
            INSERT INTO agent_commands (agent_id, command_type, payload)
            VALUES (CAST(:aid AS uuid), 'software_install', CAST(:payload AS jsonb))
        """), {"aid": agent_id, "payload": payload})
        job_ids.append(job_id)

    await db.commit()
    return {"status": "queued", "job_count": len(job_ids), "job_ids": job_ids}


@router.get("/software-deploy/jobs")
async def list_deploy_jobs(
    agent_id: Optional[str] = None,
    limit: int = 100,
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    await _ensure_tables(db)
    params = {"limit": limit}
    where = "1=1"
    if agent_id:
        where += " AND j.agent_id = CAST(:aid AS uuid)"
        params["aid"] = agent_id
    rows = await db.execute(text(f"""
        SELECT j.id, j.package_id, j.agent_id, j.status, j.output, j.error_message,
               j.triggered_by_username, j.queued_at, j.started_at, j.finished_at,
               p.name, p.version, a.hostname
        FROM software_deploy_jobs j
        LEFT JOIN software_packages p ON p.id = j.package_id
        LEFT JOIN agents a ON a.id = j.agent_id
        WHERE {where}
        ORDER BY j.queued_at DESC
        LIMIT :limit
    """), params)
    result = []
    for r in rows.fetchall():
        result.append({
            "id": str(r[0]), "package_id": str(r[1]) if r[1] else None,
            "agent_id": str(r[2]), "status": r[3], "output": r[4],
            "error_message": r[5], "triggered_by": r[6],
            "queued_at": r[7].isoformat() if r[7] else None,
            "started_at": r[8].isoformat() if r[8] else None,
            "finished_at": r[9].isoformat() if r[9] else None,
            "package_name": r[10], "package_version": r[11], "hostname": r[12],
        })
    return result


# ─── User-facing: dispatch uninstall command ─────────────────────────────────

@router.post("/agents/{agent_id}/software-uninstall")
async def dispatch_uninstall(
    agent_id: str,
    body: dict,
    db: AsyncSession = Depends(get_db),
    user=Depends(get_current_user),
):
    name    = body.get("name", "").strip()
    version = body.get("version", "") or ""
    if not name:
        raise HTTPException(status_code=400, detail="name required")
    row = await db.execute(text(
        "SELECT id FROM agents WHERE id = CAST(:id AS uuid) AND is_active = TRUE"
    ), {"id": agent_id})
    if not row.fetchone():
        raise HTTPException(status_code=404, detail="Agent not found")
    await _ensure_tables(db)
    job_id = str(uuid.uuid4())
    await db.execute(text("""
        INSERT INTO software_uninstall_jobs
            (id, agent_id, software_name, software_version, status,
             triggered_by, triggered_by_username)
        VALUES (:id, CAST(:aid AS uuid), :name, :ver, 'pending',
                CAST(:uid AS uuid), :uname)
    """), {"id": job_id, "aid": agent_id, "name": name, "ver": version,
           "uid": str(user.id), "uname": user.username})
    payload = _json.dumps({"uninstall_job_id": job_id, "name": name, "version": version})
    await db.execute(text("""
        INSERT INTO agent_commands (agent_id, command_type, payload)
        VALUES (CAST(:aid AS uuid), 'software_uninstall', CAST(:payload AS jsonb))
    """), {"aid": agent_id, "payload": payload})
    await db.commit()
    return {"status": "queued", "name": name, "job_id": job_id}


# ─── Agent-facing: uninstall result ──────────────────────────────────────────

@router.post("/agents/software-uninstall-result")
async def uninstall_result(body: dict, db: AsyncSession = Depends(get_db), agent: Agent = Depends(get_current_agent)):
    name    = body.get("name", "")
    success = body.get("success", False)
    output  = body.get("output", "")
    message = body.get("message", "")
    job_id  = body.get("uninstall_job_id", "")
    now     = datetime.now(timezone.utc)

    await _ensure_tables(db)

    # Older agent versions (< 1.4.x) don't include uninstall_job_id in the result.
    # Fall back to finding the most-recent pending job for this agent + software name.
    if not job_id and name:
        r = await db.execute(text("""
            SELECT id FROM software_uninstall_jobs
            WHERE agent_id = CAST(:aid AS uuid)
              AND software_name = :name
              AND status = 'pending'
            ORDER BY queued_at DESC
            LIMIT 1
        """), {"aid": str(agent.id), "name": name})
        row = r.fetchone()
        if row:
            job_id = str(row[0])

    if job_id:
        await db.execute(text("""
            UPDATE software_uninstall_jobs
            SET status        = :status,
                output        = :output,
                error_message = :err,
                finished_at   = :now
            WHERE id = CAST(:id AS uuid)
        """), {
            "status": "success" if success else "failed",
            "output": (output or message or "")[:10000],
            "err":    None if success else message,
            "now":    now,
            "id":     job_id,
        })

    # Also log to audit_log for general audit trail
    await db.execute(text("""
        INSERT INTO audit_log (agent_id, action, details, result)
        VALUES (CAST(:aid AS uuid), 'software_uninstall',
                CAST(:details AS jsonb), :result)
    """), {
        "aid":     str(agent.id),
        "details": _json.dumps({"name": name, "success": success,
                                 "message": message, "output": output[:2000],
                                 "job_id": job_id}),
        "result":  "success" if success else "failure",
    })
    await db.commit()
    return {"status": "ok"}


# ─── Uninstall job history ────────────────────────────────────────────────────

@router.patch("/software/uninstall-jobs/{job_id}")
async def update_uninstall_job(
    job_id: str,
    body: dict,
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    """Manually cancel or mark a stuck pending uninstall job."""
    new_status = body.get("status", "")
    if new_status not in ("cancelled", "failed", "pending"):
        raise HTTPException(400, detail="status must be cancelled, failed, or pending")
    await _ensure_tables(db)
    r = await db.execute(text(
        "UPDATE software_uninstall_jobs SET status = :s, finished_at = NOW() WHERE id = CAST(:id AS uuid)"
    ), {"s": new_status, "id": job_id})
    await db.commit()
    if r.rowcount == 0:
        raise HTTPException(404, detail="Job not found")
    return {"status": new_status}


@router.get("/software/uninstall-jobs")
async def list_uninstall_jobs(
    software_name: str = Query(""),
    agent_id:      str = Query(""),
    status:        str = Query("all"),
    limit:         int = Query(200, ge=1, le=1000),
    db: AsyncSession   = Depends(get_db),
    _=Depends(get_current_user),
):
    await _ensure_tables(db)
    conditions = []
    params: dict = {"limit": limit}
    if software_name.strip():
        conditions.append("LOWER(j.software_name) LIKE :sw")
        params["sw"] = f"%{software_name.lower()}%"
    if agent_id.strip():
        conditions.append("j.agent_id = CAST(:aid AS uuid)")
        params["aid"] = agent_id
    if status != "all":
        conditions.append("j.status = :status")
        params["status"] = status
    where = ("WHERE " + " AND ".join(conditions)) if conditions else ""
    rows = await db.execute(text(f"""
        SELECT j.id, j.software_name, j.software_version, j.status,
               j.output, j.error_message, j.triggered_by_username,
               j.queued_at, j.finished_at,
               a.hostname, a.display_name, a.ip_address, a.os_type, a.status AS agent_status
        FROM software_uninstall_jobs j
        LEFT JOIN agents a ON a.id = j.agent_id
        {where}
        ORDER BY j.queued_at DESC
        LIMIT :limit
    """), params)
    result = []
    for r in rows.fetchall():
        result.append({
            "id":              str(r[0]),
            "software_name":   r[1],
            "software_version": r[2],
            "status":          r[3],
            "output":          r[4],
            "error_message":   r[5],
            "triggered_by":    r[6],
            "queued_at":       r[7].isoformat() if r[7] else None,
            "finished_at":     r[8].isoformat() if r[8] else None,
            "hostname":        r[9],
            "display_name":    r[10],
            "ip_address":      r[11],
            "os_type":         r[12],
            "agent_status":    r[13],
        })
    return result


# ─── Agent-facing: install result ────────────────────────────────────────────

@router.post("/agents/software-install-result")
async def install_result(body: dict, db: AsyncSession = Depends(get_db), agent: Agent = Depends(get_current_agent)):
    job_id = body.get("deploy_job_id", "")
    success = body.get("success", False)
    output = body.get("output", "")
    message = body.get("message", "")
    now = datetime.now(timezone.utc)

    if job_id:
        await _ensure_tables(db)
        await db.execute(text("""
            UPDATE software_deploy_jobs
            SET status = :status, output = :output, error_message = :err,
                finished_at = :now, started_at = COALESCE(started_at, :now)
            WHERE id = CAST(:id AS uuid)
        """), {
            "status": "success" if success else "failed",
            "output": output[:10000],
            "err": None if success else message,
            "now": now,
            "id": job_id,
        })
    await db.commit()
    return {"status": "ok"}


# ─── Software Inventory (distinct list across all agents) ─────────────────────

_INVENTORY_SQL = """
    SELECT
        si.name,
        COALESCE(si.publisher, '') AS publisher,
        a.os_type,
        COUNT(DISTINCT si.agent_id)                      AS endpoint_count,
        STRING_AGG(DISTINCT si.version, ' | ')           AS versions,
        MAX(si.size_mb)                                  AS size_mb,
        MAX(si.last_seen)                                AS last_seen,
        COUNT(DISTINCT al.id)                            AS licensed_count,
        MAX(al.activation_status)                        AS lic_status,
        MAX(al.license_type)                             AS lic_type,
        MIN(al.expiry_date)                              AS lic_expiry
    FROM software_inventory si
    JOIN agents a ON a.id = si.agent_id AND a.is_active = TRUE
    LEFT JOIN agent_licenses al
           ON al.agent_id = si.agent_id
          AND LOWER(al.software_name) = LOWER(si.name)
    WHERE (
        -- For Windows agents: only show what Programs and Features shows.
        -- Exclude system components, patches, redistributables, and sub-packages.
        a.os_type != 'windows'
        OR (
            -- Exclude GUID-named entries (registry key names leaked as display names)
            si.name !~ '^\{{[0-9A-Fa-f]'
            -- Exclude Windows Update patches / hotfixes
            AND si.name !~* '^(Security Update|Update for |Hotfix for |KB[0-9]+|Cumulative Update|Service Pack|Definition Update)'
            -- Exclude all Microsoft Visual C++ Redistributables (any year/arch format)
            AND si.name !~* 'Microsoft Visual C\+\+.*[Rr]edistributable'
            AND si.name !~* 'Microsoft Visual C\+\+.*(Minimum|Additional) Runtime'
            -- Exclude .NET runtime components (not developer SDKs visible in Programs & Features)
            AND si.name !~* 'Microsoft \.NET (Framework [0-9]|Runtime [0-9]|SDK [0-9]|Targeting Pack|Windows Desktop Runtime)'
            -- Exclude Visual Studio / Windows SDK internal tools
            AND si.name !~* '^Microsoft (Visual Studio Build Tools|Visual Studio Installer)'
            AND si.name !~* '^Windows (Software Development Kit|Driver Kit|Assessment and Deployment Kit)'
            -- Exclude internal package-manager style entries (SAP/BOE sub-components etc.)
            -- These have all-lowercase dotted names with a version suffix like "foo.bar-4.0-core-32"
            AND si.name !~ '^[a-z][a-z0-9._-]*-[0-9]'
            -- Exclude Sophos internal sub-components (keep only main product entries)
            AND NOT (
                si.name ~* '^Sophos (AMSI|AutoUpdate|Diagnostic|Endpoint Defense|Endpoint Firewall|Endpoint Self Help|Exploit Prevention|File Integrity|File Scanner|Health|Lockdown|ML Engine|Management|Network Threat|Standalone|Clean)'
                OR si.name ILIKE '%managed by Sophos%'
            )
            -- Exclude Microsoft Office MUI language packs and internal sub-components.
            -- MUI = Multilingual User Interface components installed per Office app;
            -- these are never shown separately in Programs & Features.
            AND si.name !~* ' MUI \('
            AND si.name !~* 'Setup Metadata MUI'
            AND si.name !~* 'Office.*[0-9]+-bit Components'
            AND si.name !~* 'Office.*(Proofing Tools?|Proofing \(|Proof \(|OSM MUI|OSM UX MUI)'
            AND si.name !~* '^Office [0-9]+ Click-to-Run'
            AND si.name !~* 'Service Pack [0-9]'
            -- Exclude Google Update Helper (background service, not a user app)
            AND si.name !~* '^Google Update'
            -- Exclude Python installer sub-components (each Python feature registers separately)
            AND si.name !~* '^Python [0-9.]+ (Add to Path|Core Interpreter|Development Libraries|Documentation|Executables|Standard Library|Tcl.Tk Support|Test Suite|pip Bootstrap)'
            -- Exclude Visual Studio MSI sub-packages (lowercase vs_/icecap_/vcpp_ internal entries)
            AND si.name !~ '^(vs|icecap|vcpp)_'
            -- Exclude .NET SDK developer sub-components (AppHost Pack, Host FX Resolver, Toolset,
            -- Templates, Workload/SDK Manifests, Targeting Packs — not runtime installs)
            AND si.name !~* '^Microsoft \.NET (AppHost Pack|Host FX Resolver|Toolset|[0-9.]+ Templates|Standard Targeting Pack)'
            AND si.name !~* '^Microsoft\.NET\.(Sdk|Workload)\.'
            AND si.name !~* '^Microsoft (NetStandard SDK|Portable Library Multi-Targeting Pack)'
            AND si.name !~* '^Microsoft ASP\.NET Core [0-9.]+ (Targeting Pack|Hosting Bundle Options)'
            AND si.name !~* '^Microsoft Windows Desktop Targeting Pack'
            AND si.name !~* '^Microsoft \.NET Framework Cumulative Intellisense Pack'
            -- Exclude SQL Server internal sub-components (batch parsers, DMF, XEvent, CEIP, GDR patches,
            -- common/shared files, documentation, and SSIS internal sub-entries).
            -- [0-9].* matches version strings including "2008 R2 SP2", "2017", "2019" etc.
            AND si.name !~* '^(Browser for SQL Server|SQL Server Browser for SQL Server)'
            AND si.name !~* '^SQL Server [0-9].* (Batch Parser|Connection Info|DMF|XEvent|SQL Diagnostics|Shared Management Objects|SQL Data Quality Common|RS_SharePoint|Common Files|Database Engine Shared|Client Tools Extensions|Documentation Components|Integration Services (Master|Worker)|sql_)'
            AND si.name !~* '^SQL Server Integration Services Singleton'
            AND si.name !~* '^(Sql Server Customer Experience|SSMS Post Install Tasks|Prerequisites for SSDT)'
            AND si.name !~* '^(GDR|Hotfix) [0-9]+ for SQL Server'
            -- Exclude MySQL documentation and sample sub-packages
            AND si.name !~* '^MySQL (Documents|Examples and Samples)'
            -- Exclude Microsoft SQL Server internal SDK/library sub-components.
            -- [0-9][^(]* matches version strings like "2008 R2", "2012", "2019" (stops before "(64-bit)")
            AND si.name !~* '^Microsoft SQL Server [0-9][^(]*(Analysis Management Objects|Data-Tier App Framework|T-SQL Language Service|Transact-SQL|Policies|RsFx Driver|Setup |Management Objects)'
            AND si.name !~* '^Microsoft (System CLR Types for SQL Server|SQL Server System CLR Types|VSS Writer for SQL Server|AS OLE DB Provider for SQL Server)'
            AND si.name !~* '^Microsoft Visual Studio Setup (Configuration|WMI Provider)'
            -- Exclude Visual C++ runtime-only entries not covered by the Redistributable pattern
            AND si.name !~* 'Microsoft Visual C\+\+.* Runtime'
            -- Exclude Visual J# redistributable packages
            AND si.name !~* '^Microsoft Visual J#.*Redistributable'
            -- Exclude background auto-updater services (not user-facing apps)
            AND si.name !~* '^Java Auto Updater'
            AND si.name !~* '^Adobe Refresh Manager'
            -- Exclude Microsoft ASP.Net Web Frameworks security update patches (KB entries)
            AND si.name !~* '^Microsoft ASP\.Net Web Frameworks.*Security Update'
            -- Exclude SQL Server Management Studio language packs
            AND si.name !~* '^SQL Server Management Studio Language Pack'
            -- Exclude Office proofing tools in non-English UI names
            -- (Spanish "Herramientas de corrección", French "Outils de vérification linguistique")
            AND si.name !~* '^Herramientas de corrección de Microsoft Office'
            AND si.name !~* '^Outils de vérification linguistique.*de Microsoft Office'
            -- Exclude Windows SDK sub-components (developer toolchain, not end-user apps)
            AND si.name !~* '^Windows SDK (Desktop|ARM64|OnecoreUap|DirectX|Facade|Modern|Redistributabl|Signing|AddOn|EULA|for Windows Store)'
            AND si.name !~* '^Windows (App Certification Kit|Desktop Extension SDK|IoT Extension SDK|Mobile Extension SDK|Team Extension SDK)'
            AND si.name !~* '^WinRT Intellisense'
            AND si.name !~* '^Universal (CRT|General MIDI)'
            AND si.name !~* '^SDK ARM64 (Additions|Redistributables)'
            AND si.name !~* '^(WinAppDeploy|Kits Configuration Installer|Application Verifier x64 External Package|MSI Development Tools)'
            -- Exclude Bitvise internal library components (registered by installer, not separate apps)
            AND si.name !~* '^Bitvise SSH Client - FlowSshNet'
            AND si.name !~* '^WinFsp installed by'
            -- Exclude Visual Studio diagnostic and internal toolchain entries
            AND si.name !~* '^(IntelliTraceProfilerProxy|DiagnosticsHub_CollectionService|Roslyn Language Services)'
            AND si.name !~* '^VS (Immersive Activate Helper|JIT Debugger)'
            AND si.name !~* '^ClickOnce Bootstrapper Package for Microsoft'
            AND si.name !~* '^(MSI to redistribute|Visual Studio [0-9]+ Prerequisites)'
            AND si.name !~* '^Microsoft (Application Error Reporting|MPI \()'
            -- Exclude Microsoft Visual Studio TFS/Tools internal sub-components
            AND si.name !~* '^Microsoft Visual Studio Team Foundation Server.*Office Integration'
            AND si.name !~* '^Microsoft Visual Studio Tools for Applications [0-9]+ (x64|x86|Language|Finalizer)'
            AND si.name !~* '^Microsoft Visual Studio Tools for Applications (Design-Time|x64 Runtime|x86 Runtime)'
            -- Exclude Windows 10 Update Assistant (internal Windows update helper tool)
            AND si.name !~* '^Windows 10 Update Assistant'
        )
    )
    {extra_where}
    GROUP BY si.name, si.publisher, a.os_type
    ORDER BY endpoint_count DESC, si.name
"""


def _inventory_row(r) -> dict:
    return {
        "name":           r[0],
        "publisher":      r[1],
        "os_type":        r[2],
        "endpoint_count": r[3],
        "versions":       r[4] or "",
        "size_mb":        float(r[5]) if r[5] is not None else None,
        "last_seen":      r[6].isoformat() if r[6] else None,
        "licensed_count": r[7] or 0,
        "license_status": r[8],
        "license_type":   r[9],
        "license_expiry": r[10].isoformat() if r[10] else None,
    }


@router.get("/software/inventory/agents")
async def inventory_agents(
    name:    str = Query(...),
    os_type: str = Query("all"),
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    """Return all active agents that have a specific software installed."""
    conditions = ["LOWER(si.name) = LOWER(:name)", "a.is_active = TRUE"]
    params: dict = {"name": name}
    if os_type != "all":
        conditions.append("a.os_type = :os_type")
        params["os_type"] = os_type
    where = "WHERE " + " AND ".join(conditions)
    rows = await db.execute(text(f"""
        SELECT a.id, a.hostname, a.display_name, a.ip_address, a.os_type, a.status,
               STRING_AGG(DISTINCT si.version, ', ') AS versions
        FROM software_inventory si
        JOIN agents a ON a.id = si.agent_id
        {where}
        GROUP BY a.id, a.hostname, a.display_name, a.ip_address, a.os_type, a.status
        ORDER BY a.hostname
    """), params)
    return [
        {
            "id": str(r[0]), "hostname": r[1], "display_name": r[2],
            "ip_address": r[3], "os_type": r[4], "status": r[5], "versions": r[6] or "",
        }
        for r in rows.fetchall()
    ]


@router.post("/software/inventory/uninstall")
async def batch_uninstall(
    body: dict,
    db: AsyncSession = Depends(get_db),
    user=Depends(get_current_user),
):
    """
    Batch-uninstall software from multiple agents.
    Body: { "items": [{ "name": str, "version": str|null, "agent_ids": [str] }] }
    """
    items = body.get("items", [])
    if not items:
        raise HTTPException(400, detail="No items provided")

    await _ensure_tables(db)
    username = getattr(user, "username", None) or "system"
    user_id  = str(user.id)

    dispatched = 0
    errors = []
    for item in items:
        name      = (item.get("name") or "").strip()
        version   = (item.get("version") or "").strip()
        agent_ids = item.get("agent_ids") or []
        if not name or not agent_ids:
            continue
        for agent_id in agent_ids:
            try:
                r = await db.execute(text(
                    "SELECT id FROM agents WHERE id = CAST(:id AS uuid) AND is_active = TRUE"
                ), {"id": agent_id})
                if not r.fetchone():
                    errors.append(f"Agent {agent_id} not found")
                    continue
                job_id = str(uuid.uuid4())
                await db.execute(text("""
                    INSERT INTO software_uninstall_jobs
                        (id, agent_id, software_name, software_version, status,
                         triggered_by, triggered_by_username)
                    VALUES (:id, CAST(:aid AS uuid), :name, :ver, 'pending',
                            CAST(:uid AS uuid), :uname)
                """), {"id": job_id, "aid": agent_id, "name": name, "ver": version,
                       "uid": user_id, "uname": username})
                payload = _json.dumps({
                    "uninstall_job_id": job_id,
                    "name": name,
                    "version": version,
                })
                await db.execute(text("""
                    INSERT INTO agent_commands (agent_id, command_type, payload)
                    VALUES (CAST(:aid AS uuid), 'software_uninstall', CAST(:payload AS jsonb))
                """), {"aid": agent_id, "payload": payload})
                dispatched += 1
            except Exception as exc:
                errors.append(str(exc))

    if dispatched > 0:
        await db.commit()
    return {"dispatched": dispatched, "errors": errors}


@router.get("/software/inventory")
async def list_inventory(
    os_type:  str = Query("all"),
    search:   str = Query(""),
    page:     int = Query(1, ge=1),
    per_page: int = Query(100, ge=1, le=500),
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    conditions = []
    params: dict = {}
    if os_type != "all":
        conditions.append("a.os_type = :os_type")
        params["os_type"] = os_type
    if search.strip():
        conditions.append("(LOWER(si.name) LIKE :search OR LOWER(si.publisher) LIKE :search)")
        params["search"] = f"%{search.lower()}%"

    extra_where = ("AND " + " AND ".join(conditions)) if conditions else ""
    sql = _INVENTORY_SQL.format(extra_where=extra_where)

    count_sql = f"SELECT COUNT(*) FROM ({sql}) AS _sub"
    total_r = await db.execute(text(count_sql), params)
    total = total_r.scalar() or 0

    offset = (page - 1) * per_page
    paged_sql = sql + " LIMIT :limit OFFSET :offset"
    params["limit"] = per_page
    params["offset"] = offset

    rows = await db.execute(text(paged_sql), params)
    items = [_inventory_row(r) for r in rows.fetchall()]

    return {"total": total, "page": page, "per_page": per_page, "items": items}


@router.get("/software/inventory/stats")
async def inventory_stats(
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    r = await db.execute(text("""
        SELECT
            COUNT(DISTINCT si.name || '|' || COALESCE(a.os_type,''))  AS distinct_titles,
            COUNT(DISTINCT CASE WHEN a.os_type = 'windows' THEN si.name END) AS windows_titles,
            COUNT(DISTINCT CASE WHEN a.os_type = 'linux'   THEN si.name END) AS linux_titles,
            COUNT(DISTINCT al.id)                                             AS licensed_entries
        FROM software_inventory si
        JOIN agents a ON a.id = si.agent_id AND a.is_active = TRUE
        LEFT JOIN agent_licenses al ON al.agent_id = si.agent_id
               AND LOWER(al.software_name) = LOWER(si.name)
        WHERE (
            a.os_type != 'windows'
            OR (
                si.name !~ '^\{[0-9A-Fa-f]'
                AND si.name !~* '^(Security Update|Update for |Hotfix for |KB[0-9]+|Cumulative Update|Service Pack|Definition Update)'
                AND si.name !~* 'Microsoft Visual C\+\+.*[Rr]edistributable'
                AND si.name !~* 'Microsoft Visual C\+\+.*(Minimum|Additional) Runtime'
                AND si.name !~* 'Microsoft \.NET (Framework [0-9]|Runtime [0-9]|SDK [0-9]|Targeting Pack|Windows Desktop Runtime)'
                AND si.name !~* '^Microsoft (Visual Studio Build Tools|Visual Studio Installer)'
                AND si.name !~* '^Windows (Software Development Kit|Driver Kit|Assessment and Deployment Kit)'
                AND si.name !~ '^[a-z][a-z0-9._-]*-[0-9]'
                AND NOT (
                    si.name ~* '^Sophos (AMSI|AutoUpdate|Diagnostic|Endpoint Defense|Endpoint Firewall|Endpoint Self Help|Exploit Prevention|File Integrity|File Scanner|Health|Lockdown|ML Engine|Management|Network Threat|Standalone|Clean)'
                    OR si.name ILIKE '%managed by Sophos%'
                )
                AND si.name !~* ' MUI \('
                AND si.name !~* 'Setup Metadata MUI'
                AND si.name !~* 'Office.*[0-9]+-bit Components'
                AND si.name !~* 'Office.*(Proofing Tools?|Proofing \(|Proof \(|OSM MUI|OSM UX MUI)'
                AND si.name !~* '^Office [0-9]+ Click-to-Run'
                AND si.name !~* 'Service Pack [0-9]'
                AND si.name !~* '^Google Update'
                AND si.name !~* '^Python [0-9.]+ (Add to Path|Core Interpreter|Development Libraries|Documentation|Executables|Standard Library|Tcl.Tk Support|Test Suite|pip Bootstrap)'
                AND si.name !~ '^(vs|icecap|vcpp)_'
                AND si.name !~* '^Microsoft \.NET (AppHost Pack|Host FX Resolver|Toolset|[0-9.]+ Templates|Standard Targeting Pack)'
                AND si.name !~* '^Microsoft\.NET\.(Sdk|Workload)\.'
                AND si.name !~* '^Microsoft (NetStandard SDK|Portable Library Multi-Targeting Pack)'
                AND si.name !~* '^Microsoft ASP\.NET Core [0-9.]+ (Targeting Pack|Hosting Bundle Options)'
                AND si.name !~* '^Microsoft Windows Desktop Targeting Pack'
                AND si.name !~* '^Microsoft \.NET Framework Cumulative Intellisense Pack'
                AND si.name !~* '^(Browser for SQL Server|SQL Server Browser for SQL Server)'
                AND si.name !~* '^SQL Server [0-9].* (Batch Parser|Connection Info|DMF|XEvent|SQL Diagnostics|Shared Management Objects|SQL Data Quality Common|RS_SharePoint|Common Files|Database Engine Shared|Client Tools Extensions|Documentation Components|Integration Services (Master|Worker)|sql_)'
                AND si.name !~* '^SQL Server Integration Services Singleton'
                AND si.name !~* '^(Sql Server Customer Experience|SSMS Post Install Tasks|Prerequisites for SSDT)'
                AND si.name !~* '^(GDR|Hotfix) [0-9]+ for SQL Server'
                AND si.name !~* '^MySQL (Documents|Examples and Samples)'
                AND si.name !~* '^Microsoft SQL Server [0-9][^(]*(Analysis Management Objects|Data-Tier App Framework|T-SQL Language Service|Transact-SQL|Policies|RsFx Driver|Setup |Management Objects)'
                AND si.name !~* '^Microsoft (System CLR Types for SQL Server|SQL Server System CLR Types|VSS Writer for SQL Server|AS OLE DB Provider for SQL Server)'
                AND si.name !~* '^Microsoft Visual Studio Setup (Configuration|WMI Provider)'
                AND si.name !~* 'Microsoft Visual C\+\+.* Runtime'
                AND si.name !~* '^Microsoft Visual J#.*Redistributable'
                AND si.name !~* '^Java Auto Updater'
                AND si.name !~* '^Adobe Refresh Manager'
                AND si.name !~* '^Microsoft ASP\.Net Web Frameworks.*Security Update'
                AND si.name !~* '^SQL Server Management Studio Language Pack'
                AND si.name !~* '^Herramientas de corrección de Microsoft Office'
                AND si.name !~* '^Outils de vérification linguistique.*de Microsoft Office'
                AND si.name !~* '^Windows SDK (Desktop|ARM64|OnecoreUap|DirectX|Facade|Modern|Redistributabl|Signing|AddOn|EULA|for Windows Store)'
                AND si.name !~* '^Windows (App Certification Kit|Desktop Extension SDK|IoT Extension SDK|Mobile Extension SDK|Team Extension SDK)'
                AND si.name !~* '^WinRT Intellisense'
                AND si.name !~* '^Universal (CRT|General MIDI)'
                AND si.name !~* '^SDK ARM64 (Additions|Redistributables)'
                AND si.name !~* '^(WinAppDeploy|Kits Configuration Installer|Application Verifier x64 External Package|MSI Development Tools)'
                AND si.name !~* '^Bitvise SSH Client - FlowSshNet'
                AND si.name !~* '^WinFsp installed by'
                AND si.name !~* '^(IntelliTraceProfilerProxy|DiagnosticsHub_CollectionService|Roslyn Language Services)'
                AND si.name !~* '^VS (Immersive Activate Helper|JIT Debugger)'
                AND si.name !~* '^ClickOnce Bootstrapper Package for Microsoft'
                AND si.name !~* '^(MSI to redistribute|Visual Studio [0-9]+ Prerequisites)'
                AND si.name !~* '^Microsoft (Application Error Reporting|MPI \()'
                AND si.name !~* '^Microsoft Visual Studio Team Foundation Server.*Office Integration'
                AND si.name !~* '^Microsoft Visual Studio Tools for Applications [0-9]+ (x64|x86|Language|Finalizer)'
                AND si.name !~* '^Microsoft Visual Studio Tools for Applications (Design-Time|x64 Runtime|x86 Runtime)'
                AND si.name !~* '^Windows 10 Update Assistant'
            )
        )
    """))
    row = r.fetchone()
    return {
        "distinct_titles":  row[0] or 0,
        "windows_titles":   row[1] or 0,
        "linux_titles":     row[2] or 0,
        "licensed_entries": row[3] or 0,
    }


@router.get("/software/inventory/export")
async def export_inventory_xlsx(
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(500, detail="openpyxl not installed")

    rows_w = (await db.execute(text(_INVENTORY_SQL.format(extra_where="AND a.os_type = :os_type")), {"os_type": "windows"})).fetchall()
    rows_l = (await db.execute(text(_INVENTORY_SQL.format(extra_where="AND a.os_type = :os_type")), {"os_type": "linux"})).fetchall()

    DARK_BG     = "0B1120"; HEADER_BLUE = "1D4ED8"; SECTION_BG = "1E293B"
    WHITE       = "FFFFFF"; LIGHT_SLATE = "CBD5E1"
    GREEN_BG    = "D1FAE5"; GREEN_FG = "065F46"
    AMBER_BG    = "FEF3C7"; AMBER_FG = "92400E"
    RED_BG      = "FEE2E2"; RED_FG   = "991B1B"

    def _fill(h): return PatternFill("solid", fgColor=h)
    def _font(bold=False, color=WHITE, size=10): return Font(bold=bold, color=color, size=size, name="Calibri")
    def _border():
        s = Side(style="thin", color="334155")
        return Border(left=s, right=s, top=s, bottom=s)
    def _center(): return Alignment(horizontal="center", vertical="center", wrap_text=True)
    def _left():   return Alignment(horizontal="left",   vertical="center", wrap_text=True)

    HEADERS = ["Software Name", "Publisher", "Version(s)", "Endpoints",
               "Size (MB)", "Last Seen", "Licensed Seats",
               "License Status", "License Type", "Expiry Date"]
    COL_W   = [42, 32, 24, 12, 12, 18, 15, 18, 18, 16]

    def _build_sheet(ws, title: str, data_rows):
        ws.merge_cells(f"A1:{get_column_letter(len(HEADERS))}1")
        tc = ws["A1"]
        tc.value = title
        tc.fill = _fill(DARK_BG); tc.font = _font(bold=True, size=13)
        tc.alignment = _center(); ws.row_dimensions[1].height = 30

        for i, w in enumerate(COL_W, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        for col, hdr in enumerate(HEADERS, 1):
            c = ws.cell(row=2, column=col, value=hdr)
            c.fill = _fill(HEADER_BLUE); c.font = _font(bold=True)
            c.alignment = _center(); c.border = _border()
        ws.row_dimensions[2].height = 22
        ws.freeze_panes = "A3"

        for ri, r in enumerate(data_rows, 3):
            d = _inventory_row(r)
            vals = [
                d["name"], d["publisher"], d["versions"],
                d["endpoint_count"],
                round(d["size_mb"], 1) if d["size_mb"] else "",
                d["last_seen"][:10] if d["last_seen"] else "",
                d["licensed_count"] or "",
                d["license_status"] or "",
                d["license_type"] or "",
                d["license_expiry"][:10] if d["license_expiry"] else "",
            ]
            for ci, val in enumerate(vals, 1):
                c = ws.cell(row=ri, column=ci, value=val)
                c.fill = _fill(SECTION_BG); c.font = _font(color=LIGHT_SLATE)
                c.alignment = _left() if ci <= 3 else _center()
                c.border = _border()

            lic_cell = ws.cell(row=ri, column=8)
            st = (d["license_status"] or "").lower()
            if st in ("licensed", "activated", "active"):
                lic_cell.fill = _fill(GREEN_BG); lic_cell.font = _font(bold=True, color=GREEN_FG)
            elif st in ("trial", "expiring"):
                lic_cell.fill = _fill(AMBER_BG); lic_cell.font = _font(bold=True, color=AMBER_FG)
            elif st in ("unlicensed", "expired", "invalid"):
                lic_cell.fill = _fill(RED_BG); lic_cell.font = _font(bold=True, color=RED_FG)
            lic_cell.alignment = _center()
            ws.row_dimensions[ri].height = 18

    wb = openpyxl.Workbook()
    ws_win = wb.active
    ws_win.title = "Windows Software"
    _build_sheet(ws_win, f"Windows Software Inventory  ({len(rows_w)} distinct titles)", rows_w)

    ws_lin = wb.create_sheet("Linux Software")
    _build_sheet(ws_lin, f"Linux Software Inventory  ({len(rows_l)} distinct titles)", rows_l)

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    now_str = datetime.now().strftime("%Y%m%d_%H%M")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="software_inventory_{now_str}.xlsx"'},
    )
