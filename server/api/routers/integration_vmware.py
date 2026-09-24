"""VMware vCenter integration data router."""
from fastapi import APIRouter, Depends, HTTPException, Query
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from api.database import get_db
from api.services.auth import get_current_user, require_operator

router = APIRouter(prefix="/integrations/vmware", tags=["Integrations - VMware"])


@router.get("/dashboard")
async def vmware_dashboard(db: AsyncSession = Depends(get_db), _=Depends(get_current_user)):
    plugin = await db.execute(text(
        "SELECT status, is_enabled, last_sync_at, last_error FROM integration_plugins WHERE plugin_type='vmware'"
    ))
    p = plugin.fetchone()

    vms = await db.execute(text("""
        SELECT COUNT(*) AS total,
               COUNT(*) FILTER (WHERE power_state = 'POWERED_ON') AS powered_on,
               COUNT(*) FILTER (WHERE power_state = 'POWERED_OFF') AS powered_off,
               COUNT(*) FILTER (WHERE power_state = 'SUSPENDED') AS suspended,
               SUM(cpu_count) AS total_vcpus,
               SUM(memory_mb) AS total_memory_mb
        FROM vmware_vms
    """))
    v = vms.fetchone()

    hosts = await db.execute(text("""
        SELECT COUNT(*) AS total,
               COUNT(*) FILTER (WHERE connection_state = 'connected') AS connected,
               SUM(cpu_cores) AS total_cores,
               SUM(memory_mb) AS total_memory_mb,
               SUM(cpu_usage_mhz) AS total_cpu_usage_mhz,
               SUM(cpu_cores * cpu_mhz) AS total_cpu_capacity_mhz,
               SUM(memory_usage_mb) AS total_mem_usage_mb
        FROM vmware_hosts
    """))
    h = hosts.fetchone()

    ds = await db.execute(text("""
        SELECT COUNT(*) AS total,
               SUM(capacity_mb) AS total_mb,
               SUM(free_mb) AS free_mb
        FROM vmware_datastores
    """))
    d = ds.fetchone()

    clusters = await db.execute(text("SELECT COUNT(*) AS total FROM vmware_clusters"))
    cl = clusters.fetchone()

    cpu_pct = 0
    mem_pct = 0
    if h.total_cpu_capacity_mhz and h.total_cpu_capacity_mhz > 0:
        cpu_pct = round((h.total_cpu_usage_mhz or 0) / h.total_cpu_capacity_mhz * 100, 1)
    if h.total_memory_mb and h.total_memory_mb > 0:
        mem_pct = round((h.total_mem_usage_mb or 0) / h.total_memory_mb * 100, 1)

    ds_used_mb = (d.total_mb or 0) - (d.free_mb or 0)
    ds_pct = round(ds_used_mb / d.total_mb * 100, 1) if d.total_mb else 0

    return {
        "plugin": {
            "status": p.status if p else "disconnected",
            "is_enabled": p.is_enabled if p else False,
            "last_sync_at": p.last_sync_at.isoformat() if p and p.last_sync_at else None,
            "last_error": p.last_error if p else None,
        },
        "vms": {
            "total": v.total or 0,
            "powered_on": v.powered_on or 0,
            "powered_off": v.powered_off or 0,
            "suspended": v.suspended or 0,
            "total_vcpus": v.total_vcpus or 0,
            "total_memory_mb": v.total_memory_mb or 0,
        },
        "hosts": {
            "total": h.total or 0,
            "connected": h.connected or 0,
            "total_cores": h.total_cores or 0,
            "cpu_usage_pct": cpu_pct,
            "memory_usage_pct": mem_pct,
        },
        "datastores": {
            "total": d.total or 0,
            "total_mb": d.total_mb or 0,
            "free_mb": d.free_mb or 0,
            "usage_pct": ds_pct,
        },
        "clusters": {"total": cl.total or 0},
    }


@router.get("/vms")
async def list_vms(
    power_state: str = None,
    search: str = None,
    limit: int = Query(default=200, ge=1, le=1000),
    db: AsyncSession = Depends(get_db), _=Depends(get_current_user),
):
    filters, params = [], {"limit": limit}
    if power_state:
        filters.append("power_state = :ps")
        params["ps"] = power_state.upper()
    if search:
        filters.append("(name ILIKE :q OR guest_os ILIKE :q OR ip_address ILIKE :q OR host_name ILIKE :q)")
        params["q"] = f"%{search}%"
    where = ("WHERE " + " AND ".join(filters)) if filters else ""
    rows = await db.execute(text(f"""
        SELECT vm_id, name, power_state, cpu_count, memory_mb, guest_os,
               ip_address, host_name, cluster_name, cpu_usage_mhz,
               memory_usage_mb, tools_status, vcenter_created_at
        FROM vmware_vms {where}
        ORDER BY power_state, name LIMIT :limit
    """), params)
    return [dict(r._mapping) for r in rows.fetchall()]


@router.get("/vms/export")
async def export_vms(db: AsyncSession = Depends(get_db), _=Depends(get_current_user)):
    """Export all VMs as XLSX."""
    import io
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")

    rows = await db.execute(text("""
        SELECT name, power_state, guest_os, cpu_count, memory_mb,
               ip_address, host_name, cluster_name, tools_status,
               vcenter_created_at, synced_at
        FROM vmware_vms ORDER BY power_state, name
    """))
    vms = rows.fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VMs"

    hdr_fill = PatternFill("solid", fgColor="1E293B")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    headers = ["Name", "Power State", "Guest OS", "CPUs", "Memory",
               "IP Address", "Host", "Cluster", "Tools Status", "Date Created", "Last Synced"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = hdr_font; c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center")

    def fmt_mb(mb):
        if not mb: return "0 GB"
        if mb >= 1024 * 1024: return f"{mb/1024/1024:.1f} TB"
        if mb >= 1024: return f"{mb/1024:.1f} GB"
        return f"{mb} MB"

    for row_i, vm in enumerate(vms, 2):
        vals = [
            vm[0], vm[1], vm[2] or "—", vm[3] or 0, fmt_mb(vm[4]),
            vm[5] or "—", vm[6] or "—", vm[7] or "—", vm[8] or "—",
            vm[9].strftime("%Y-%m-%d %H:%M") if vm[9] else "—",
            vm[10].strftime("%Y-%m-%d %H:%M") if vm[10] else "—",
        ]
        row_fill = PatternFill("solid", fgColor="1F2937" if row_i % 2 == 0 else "111827")
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=row_i, column=col, value=val)
            c.font = Font(color="E2E8F0", size=9)
            c.fill = row_fill
            c.alignment = Alignment(horizontal="left")
        # Color power state cell
        ps_cell = ws.cell(row=row_i, column=2)
        if vm[1] == "POWERED_ON":
            ps_cell.font = Font(bold=True, color="34D399", size=9)
        elif vm[1] == "POWERED_OFF":
            ps_cell.font = Font(color="94A3B8", size=9)

    col_widths = [28, 14, 28, 6, 10, 16, 22, 18, 14, 18, 18]
    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    from fastapi.responses import Response
    from datetime import date
    filename = f"vmware_vms_{date.today()}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/hosts")
async def list_hosts(db: AsyncSession = Depends(get_db), _=Depends(get_current_user)):
    rows = await db.execute(text("""
        SELECT host_id, name, connection_state, power_state,
               cpu_cores, cpu_mhz, cpu_usage_mhz, memory_mb, memory_usage_mb,
               vm_count, cluster_name, version,
               CASE WHEN cpu_cores * cpu_mhz > 0
                    THEN ROUND(cpu_usage_mhz::numeric / (cpu_cores * cpu_mhz) * 100, 1)
                    ELSE 0 END AS cpu_usage_pct,
               CASE WHEN memory_mb > 0
                    THEN ROUND(memory_usage_mb::numeric / memory_mb * 100, 1)
                    ELSE 0 END AS memory_usage_pct
        FROM vmware_hosts ORDER BY name
    """))
    return [dict(r._mapping) for r in rows.fetchall()]


@router.get("/datastores")
async def list_datastores(db: AsyncSession = Depends(get_db), _=Depends(get_current_user)):
    rows = await db.execute(text("""
        SELECT ds_id, name, ds_type, capacity_mb, free_mb, accessible,
               CASE WHEN capacity_mb > 0
                    THEN ROUND((capacity_mb - free_mb)::numeric / capacity_mb * 100, 1)
                    ELSE 0 END AS usage_pct
        FROM vmware_datastores ORDER BY name
    """))
    return [dict(r._mapping) for r in rows.fetchall()]


@router.get("/clusters")
async def list_clusters(db: AsyncSession = Depends(get_db), _=Depends(get_current_user)):
    rows = await db.execute(text("""
        SELECT cluster_id, name, ha_enabled, drs_enabled, host_count, vm_count, cpu_cores, memory_mb
        FROM vmware_clusters ORDER BY name
    """))
    return [dict(r._mapping) for r in rows.fetchall()]


@router.get("/agent-vm/{agent_id}")
async def get_agent_vmware_vm(
    agent_id: str,
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    """Return the VMware VM that matches this agent by IP or hostname (for power-on button)."""
    row = await db.execute(text(
        "SELECT id, ip_address, hostname FROM agents WHERE id = CAST(:aid AS uuid) AND is_active = TRUE"
    ), {"aid": agent_id})
    agent = row.fetchone()
    if not agent:
        raise HTTPException(404, "Agent not found")

    agent_ip = agent[1] or ""
    agent_host = (agent[2] or "").lower()

    vm_row = await db.execute(text("""
        SELECT name, power_state, ip_address, guest_os
        FROM vmware_vms
        WHERE LOWER(name) = :host
           OR ip_address = :ip
        LIMIT 1
    """), {"host": agent_host, "ip": agent_ip})
    vm = vm_row.fetchone()
    if not vm:
        return {"found": False}
    return {
        "found": True,
        "vm_name": vm[0],
        "power_state": vm[1],
        "vm_ip": vm[2],
        "guest_os": vm[3],
    }


@router.post("/power-on")
async def vmware_power_on(
    body: dict,
    db: AsyncSession = Depends(get_db),
    _=Depends(require_operator),
):
    """Power on a VMware VM by name. Used to start an offline agent's host VM."""
    import asyncio

    vm_name = (body.get("vm_name") or "").strip()
    if not vm_name:
        raise HTTPException(400, "vm_name is required")

    plugin = await db.execute(text(
        "SELECT config, is_enabled FROM integration_plugins WHERE plugin_type = 'vmware'"
    ))
    plugin = plugin.fetchone()
    if not plugin or not plugin[1]:
        raise HTTPException(503, "VMware integration is not enabled")

    cfg = plugin[0] or {}

    def _power_on():
        from api.services.vmware_client import VMwareClient
        sources = []
        if cfg.get("host"):
            sources.append(VMwareClient(
                host=cfg["host"],
                username=cfg.get("username", ""),
                password=cfg.get("password", ""),
                verify_ssl=cfg.get("verify_ssl", False),
            ))
        for esxi in (cfg.get("standalone_esxi") or []):
            if esxi.get("host"):
                sources.append(VMwareClient(
                    host=esxi["host"],
                    username=esxi.get("username", ""),
                    password=esxi.get("password", ""),
                    verify_ssl=esxi.get("verify_ssl", False),
                    is_esxi=True,
                ))
        for client in sources:
            try:
                client.login()
                if client.power_on_vm(vm_name):
                    client.close()
                    return True
                client.close()
            except Exception:
                try:
                    client.close()
                except Exception:
                    pass
        return False

    found = await asyncio.get_event_loop().run_in_executor(None, _power_on)
    if not found:
        raise HTTPException(404, f"VM '{vm_name}' not found on any VMware source")
    return {"status": "power_on_task_started", "vm_name": vm_name}
