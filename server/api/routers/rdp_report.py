"""RDP Connection Report — per-server session data exported as XLSX."""
import io
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from api.database import get_db
from api.services.auth import get_current_user

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

router = APIRouter(prefix="/rdp-report", tags=["rdp-report"])

# ── Styles ────────────────────────────────────────────────────────────────────
_HDR_FILL  = PatternFill("solid", fgColor="1F3864") if HAS_OPENPYXL else None
_HDR_FONT  = Font(bold=True, color="FFFFFF", size=11) if HAS_OPENPYXL else None
_ALT_FILL  = PatternFill("solid", fgColor="EAF0FB") if HAS_OPENPYXL else None
_BOLD      = Font(bold=True) if HAS_OPENPYXL else None
_CENTER    = Alignment(horizontal="center", vertical="center") if HAS_OPENPYXL else None
_WRAP      = Alignment(wrap_text=True, vertical="top") if HAS_OPENPYXL else None
_THIN      = Side(border_style="thin", color="CCCCCC") if HAS_OPENPYXL else None
_BORDER    = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN) if HAS_OPENPYXL else None

_SUM_FILL  = PatternFill("solid", fgColor="2E75B6") if HAS_OPENPYXL else None
_SUM_FONT  = Font(bold=True, color="FFFFFF", size=12) if HAS_OPENPYXL else None

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def _hdr(ws, row, col, value, fill=None, font=None, align=None, border=None):
    cell = ws.cell(row=row, column=col, value=value)
    if fill:  cell.fill   = fill
    if font:  cell.font   = font
    if align: cell.alignment = align
    if border: cell.border = border
    return cell


def _fmt_duration(seconds):
    if seconds is None:
        return "—"
    h = seconds // 3600
    m = (seconds % 3600) // 60
    s = seconds % 60
    if h:
        return f"{h}h {m:02d}m"
    if m:
        return f"{m}m {s:02d}s"
    return f"{s}s"


def _autowidth(ws, min_w=8, max_w=50):
    for col in ws.columns:
        best = min_w
        for cell in col:
            try:
                best = max(best, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(best + 2, max_w)


# ── Endpoints ─────────────────────────────────────────────────────────────────

@router.get("/agents")
async def list_rdp_agents(
    db: AsyncSession = Depends(get_db),
    _user=Depends(get_current_user),
):
    """List all Windows agents for the server-picker UI, with RDP session counts."""
    rows = await db.execute(text("""
        SELECT a.id, a.hostname, a.display_name, a.ip_address, a.os_type, a.status,
               COUNT(r.id) AS session_count,
               MIN(r.logon_time) AS earliest,
               MAX(r.logon_time) AS latest,
               COUNT(r.id) FILTER (
                   WHERE r.logoff_time IS NULL
                     AND r.logon_time > NOW() - INTERVAL '7 days'
               ) AS active_count
        FROM agents a
        LEFT JOIN rdp_sessions r ON r.agent_id = a.id
        WHERE a.is_active = TRUE AND a.os_type = 'windows'
        GROUP BY a.id, a.hostname, a.display_name, a.ip_address, a.os_type, a.status
        ORDER BY a.hostname
    """))
    return [
        {
            "id":            str(r[0]),
            "hostname":      r[1],
            "display_name":  r[2] or r[1],
            "ip_address":    r[3],
            "os_type":       r[4],
            "status":        r[5],
            "session_count": r[6],
            "earliest":      r[7].isoformat() if r[7] else None,
            "latest":        r[8].isoformat() if r[8] else None,
            "active_count":  r[9],
        }
        for r in rows.fetchall()
    ]


@router.get("/sessions/{agent_id}")
async def get_agent_sessions(
    agent_id: str,
    year: int = None,
    month: int = None,
    db: AsyncSession = Depends(get_db),
    _user=Depends(get_current_user),
):
    """Return recent RDP sessions for one agent, optionally filtered by year/month."""
    from datetime import datetime, timezone
    filters = "r.agent_id = CAST(:aid AS uuid)"
    params = {"aid": agent_id}
    if year:
        filters += " AND EXTRACT(YEAR FROM r.logon_time AT TIME ZONE 'UTC') = :yr"
        params["yr"] = year
    if month:
        filters += " AND EXTRACT(MONTH FROM r.logon_time AT TIME ZONE 'UTC') = :mo"
        params["mo"] = month

    rows = await db.execute(text(f"""
        SELECT r.username, r.domain, r.source_ip, r.logon_time,
               r.logoff_time, r.duration_seconds, r.logoff_type,
               a.hostname, a.display_name, a.ip_address, a.status
        FROM rdp_sessions r
        JOIN agents a ON a.id = r.agent_id
        WHERE {filters}
        ORDER BY r.logon_time DESC
        LIMIT 500
    """), params)

    sessions = []
    for r in rows.fetchall():
        sessions.append({
            "username":         r[0],
            "domain":           r[1],
            "source_ip":        r[2],
            "logon_time":       r[3].isoformat() if r[3] else None,
            "logoff_time":      r[4].isoformat() if r[4] else None,
            "duration_seconds": r[5],
            "logoff_type":      r[6],
            "hostname":         r[7],
            "display_name":     r[8] or r[7],
            "ip_address":       r[9],
            "server_status":    r[10],
        })
    return {"agent_id": agent_id, "sessions": sessions}


@router.get("/active/{agent_id}")
async def get_active_sessions(
    agent_id: str,
    db: AsyncSession = Depends(get_db),
    _user=Depends(get_current_user),
):
    """Return sessions with no logoff (currently connected or recently disconnected)."""
    rows = await db.execute(text("""
        SELECT r.username, r.domain, r.source_ip, r.logon_time,
               r.logoff_type, a.hostname, a.display_name, a.ip_address
        FROM rdp_sessions r
        JOIN agents a ON a.id = r.agent_id
        WHERE r.agent_id = CAST(:aid AS uuid)
          AND r.logoff_time IS NULL
          AND r.logon_time > NOW() - INTERVAL '7 days'
        ORDER BY r.logon_time DESC
    """), {"aid": agent_id})

    return [
        {
            "username":     r[0],
            "domain":       r[1],
            "source_ip":    r[2],
            "logon_time":   r[3].isoformat() if r[3] else None,
            "logoff_type":  r[4],
            "hostname":     r[5],
            "display_name": r[6] or r[5],
            "ip_address":   r[7],
        }
        for r in rows.fetchall()
    ]


@router.get("/xlsx")
async def rdp_report_xlsx(
    agent_ids: str = Query(..., description="Comma-separated agent UUIDs"),
    year: int = Query(default=None, description="Year to report (default: current year)"),
    db: AsyncSession = Depends(get_db),
    _user=Depends(get_current_user),
):
    """Generate XLSX RDP connection report for selected servers."""
    if not HAS_OPENPYXL:
        raise HTTPException(status_code=500, detail="openpyxl not installed")

    report_year = year or datetime.now(timezone.utc).year
    ids = [a.strip() for a in agent_ids.split(",") if a.strip()]
    if not ids:
        raise HTTPException(status_code=400, detail="No agent IDs provided")

    # ── Fetch agent info ──────────────────────────────────────────────────────
    agents_rows = (await db.execute(text("""
        SELECT id, hostname, display_name, ip_address
        FROM agents WHERE id = ANY(CAST(:ids AS uuid[]))
        ORDER BY hostname
    """), {"ids": ids})).fetchall()

    if not agents_rows:
        raise HTTPException(status_code=404, detail="No matching agents found")

    agent_map = {str(r[0]): {"hostname": r[1], "display_name": r[2] or r[1], "ip": r[3]}
                 for r in agents_rows}

    # ── Fetch sessions for the year ───────────────────────────────────────────
    sessions_rows = (await db.execute(text("""
        SELECT r.agent_id, r.username, r.domain, r.source_ip,
               r.logon_time, r.logoff_time, r.duration_seconds, r.logoff_type
        FROM rdp_sessions r
        WHERE r.agent_id = ANY(CAST(:ids AS uuid[]))
          AND EXTRACT(YEAR FROM r.logon_time AT TIME ZONE 'UTC') = :yr
        ORDER BY r.agent_id, r.logon_time
    """), {"ids": ids, "yr": report_year})).fetchall()

    # ── Group sessions by agent ───────────────────────────────────────────────
    by_agent: dict[str, list] = {aid: [] for aid in ids}
    for row in sessions_rows:
        aid = str(row[0])
        if aid in by_agent:
            by_agent[aid].append(row)

    # ── Summary: connections per agent per month ──────────────────────────────
    # summary[agent_id][month_index(0-11)] = count
    summary: dict[str, list[int]] = {aid: [0]*12 for aid in ids}
    for row in sessions_rows:
        aid = str(row[0])
        month_idx = row[4].month - 1  # logon_time.month - 1
        if aid in summary:
            summary[aid][month_idx] += 1

    # ── Build workbook ────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    # ─── Sheet 1: Monthly Summary ─────────────────────────────────────────────
    ws_sum = wb.create_sheet("Monthly Summary")
    ws_sum.freeze_panes = "B3"

    # Title
    ws_sum.merge_cells("A1:N1")
    t = ws_sum["A1"]
    t.value = f"RDP Connection Report — {report_year}"
    t.font = Font(bold=True, size=14, color="FFFFFF")
    t.fill = _SUM_FILL
    t.alignment = _CENTER
    ws_sum.row_dimensions[1].height = 28

    # Header row — Server | Jan | Feb | ... | Dec | Total
    headers = ["Server"] + MONTHS + ["Total"]
    for ci, h in enumerate(headers, 1):
        _hdr(ws_sum, 2, ci, h, fill=_HDR_FILL, font=_HDR_FONT, align=_CENTER, border=_BORDER)
    ws_sum.row_dimensions[2].height = 22

    total_by_month = [0] * 12
    grand_total = 0
    row_idx = 3
    for rid, (aid, info) in enumerate(
        sorted(agent_map.items(), key=lambda x: x[1]["hostname"])
    ):
        counts = summary[aid]
        row_total = sum(counts)
        fill = _ALT_FILL if rid % 2 == 1 else None

        ws_sum.cell(row=row_idx, column=1, value=info["display_name"]).border = _BORDER
        if fill:
            ws_sum.cell(row=row_idx, column=1).fill = fill

        for mi, cnt in enumerate(counts):
            c = ws_sum.cell(row=row_idx, column=mi + 2, value=cnt if cnt else "")
            c.alignment = _CENTER
            c.border = _BORDER
            if fill:
                c.fill = fill
            total_by_month[mi] += cnt

        tc = ws_sum.cell(row=row_idx, column=14, value=row_total)
        tc.font = _BOLD
        tc.alignment = _CENTER
        tc.border = _BORDER
        if fill:
            tc.fill = fill
        grand_total += row_total
        row_idx += 1

    # Totals row
    tot_row = row_idx
    ws_sum.cell(tot_row, 1, "TOTAL").font = _BOLD
    ws_sum.cell(tot_row, 1).fill = PatternFill("solid", fgColor="D6E4F0")
    ws_sum.cell(tot_row, 1).border = _BORDER
    for mi, cnt in enumerate(total_by_month):
        c = ws_sum.cell(tot_row, mi + 2, cnt if cnt else "")
        c.font = _BOLD
        c.alignment = _CENTER
        c.fill = PatternFill("solid", fgColor="D6E4F0")
        c.border = _BORDER
    gc = ws_sum.cell(tot_row, 14, grand_total)
    gc.font = Font(bold=True, size=11)
    gc.alignment = _CENTER
    gc.fill = PatternFill("solid", fgColor="BDD7EE")
    gc.border = _BORDER

    ws_sum.column_dimensions["A"].width = 28
    for ci in range(2, 15):
        ws_sum.column_dimensions[get_column_letter(ci)].width = 8

    # ─── Per-server sheets ────────────────────────────────────────────────────
    HDR = ["#", "Username", "Domain", "Source IP", "Logon Time", "Logoff Time",
           "Duration", "Session End"]
    HDR_WIDTHS = [5, 22, 18, 18, 22, 22, 12, 14]

    for aid, info in sorted(agent_map.items(), key=lambda x: x[1]["hostname"]):
        sheet_name = info["hostname"][:31]  # Excel tab limit
        ws = wb.create_sheet(sheet_name)
        ws.freeze_panes = "A3"

        # Title
        ws.merge_cells(f"A1:{get_column_letter(len(HDR))}1")
        t = ws["A1"]
        t.value = f"{info['display_name']}  ({info['ip']})  —  RDP Sessions {report_year}"
        t.font = Font(bold=True, size=13, color="FFFFFF")
        t.fill = _SUM_FILL
        t.alignment = _CENTER
        ws.row_dimensions[1].height = 26

        # Column headers
        for ci, (h, w) in enumerate(zip(HDR, HDR_WIDTHS), 1):
            _hdr(ws, 2, ci, h, fill=_HDR_FILL, font=_HDR_FONT, align=_CENTER, border=_BORDER)
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.row_dimensions[2].height = 20

        sessions = by_agent[aid]
        if not sessions:
            ws.merge_cells(f"A3:{get_column_letter(len(HDR))}3")
            ws["A3"].value = "No RDP sessions recorded for this server in the selected year."
            ws["A3"].alignment = _CENTER
            continue

        for ri, row in enumerate(sessions):
            r = ri + 3
            fill = _ALT_FILL if ri % 2 == 1 else None
            logon_dt  = row[4]
            logoff_dt = row[5]
            dur_secs  = row[6]
            logoff_type = row[7] or "unknown"

            logon_str  = logon_dt.strftime("%Y-%m-%d %H:%M:%S") if logon_dt else "—"
            logoff_str = logoff_dt.strftime("%Y-%m-%d %H:%M:%S") if logoff_dt else "—"

            values = [
                ri + 1,
                row[1],        # username
                row[2] or "",  # domain
                row[3] or "",  # source_ip
                logon_str,
                logoff_str,
                _fmt_duration(dur_secs),
                logoff_type.title(),
            ]
            for ci, val in enumerate(values, 1):
                c = ws.cell(r, ci, val)
                c.border = _BORDER
                c.alignment = _CENTER if ci in (1, 7, 8) else Alignment(vertical="center")
                if fill:
                    c.fill = fill

    # ── Stream response ───────────────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"rdp_report_{report_year}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
