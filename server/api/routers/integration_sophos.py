"""
Sophos Central integration data router.
Returns cached data from sophos_endpoints and sophos_alerts tables.
"""
import re
from datetime import datetime, timezone, date as _date
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query, Request
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from api.database import get_db
from api.services.auth import get_current_user


# ── Incident report derivation ───────────────────────────────────────────────

def _derive_incident(category: str, description: str, endpoint_hostname: str) -> dict:
    """Derive structured incident fields from raw alert data."""
    desc = description or ""
    desc_l = desc.lower()
    cat = (category or "").lower()

    # ── Auto-resolved detection ──────────────────────────────────────────────
    auto_kw = ["terminated successfully", "blocked", "quarantined", "cleaned successfully",
                "removed successfully", "accesspoint.online", "event.wireless.accesspoint.online"]
    manual_kw = ["reboot required", "is down", "is offline", "suspended",
                 "botnet", "command and control", "not reporting", "manual"]
    auto_resolved = any(kw in desc_l for kw in auto_kw) and not any(kw in desc_l for kw in manual_kw)

    # ── Incident Type ────────────────────────────────────────────────────────
    type_map = {
        "connectivity":       "Network / VPN Connectivity",
        "malware":            "Malware Detection",
        "pua":                "Potentially Unwanted Application (PUA)",
        "runtimedetections":  "Runtime Threat Detection",
        "security":           "Security Threat",
        "protection":         "Endpoint Protection",
        "wireless":           "Wireless Connectivity",
        "updating":           "Software Update",
        "general":            "General Alert",
    }
    incident_type = type_map.get(cat, (category or "Unknown").replace("_", " ").title())

    # ── Source of the issue ──────────────────────────────────────────────────
    source_of_issue = endpoint_hostname or ""
    if cat == "wireless":
        m = re.search(r'access point[:\s"\']*([A-Z0-9_\-]+(?:\s+[^\[,\n]+)?)', desc, re.IGNORECASE)
        if m:
            source_of_issue = m.group(1).strip(" '\"[]")
        elif not source_of_issue:
            source_of_issue = "Wireless Infrastructure"
    if not source_of_issue:
        source_of_issue = "Unknown"

    # ── Affected System/User ─────────────────────────────────────────────────
    affected = endpoint_hostname or ""
    if cat == "connectivity":
        m = re.search(r"user\s+(.+?)\s+was\s+(terminated|disconnected)", desc, re.IGNORECASE)
        if m:
            user = m.group(1).strip()
            affected = f"{endpoint_hostname} — User: {user}" if endpoint_hostname else user
    if not affected:
        affected = source_of_issue

    # ── Initial Response ─────────────────────────────────────────────────────
    if cat == "connectivity":
        if "terminated successfully" in desc_l:
            initial_response = "VPN session terminated automatically by Sophos Central"
        elif "is down" in desc_l:
            initial_response = "Gateway connectivity loss detected — alert raised; manual response required"
        else:
            initial_response = "Connectivity event detected by Sophos Central"
    elif cat == "malware":
        initial_response = "Malware detected; quarantine/removal attempted by Sophos"
    elif cat == "pua":
        if "reboot required" in desc_l:
            initial_response = "PUA detected and flagged — system reboot required to complete removal"
        else:
            initial_response = "PUA detected and removed by Sophos"
    elif cat == "security":
        initial_response = "Threat detected; communication blocked by Sophos — manual investigation required"
    elif cat == "protection":
        initial_response = "Protection service disruption detected — manual response required"
    elif cat == "wireless":
        if "offline" in desc_l or "is offline" in desc_l:
            initial_response = "Access Point offline event detected — manual investigation required"
        elif "online" in desc_l:
            initial_response = "Access Point connectivity restored — auto-detected by Sophos"
        else:
            initial_response = "Wireless event detected by Sophos Central"
    elif cat == "runtimedetections":
        initial_response = "Runtime threat detected and blocked by Sophos"
    elif cat == "updating":
        initial_response = "Software update event detected by Sophos"
    else:
        initial_response = "Alert raised by Sophos Central; under review"

    # ── Resolution ───────────────────────────────────────────────────────────
    if auto_resolved:
        if cat == "connectivity" and "terminated" in desc_l:
            resolution = "Automatically resolved — VPN session terminated successfully by Sophos"
        elif cat == "wireless" and "online" in desc_l:
            resolution = "Automatically resolved — wireless access point restored"
        else:
            resolution = "Automatically resolved by Sophos Central"
    elif "reboot required" in desc_l:
        resolution = "Pending — system reboot required to complete remediation"
    elif cat == "security":
        resolution = "Pending — security investigation and manual remediation required"
    elif cat == "protection":
        resolution = "Pending — protection service requires manual restart/investigation"
    elif cat == "wireless" and "offline" in desc_l:
        resolution = "Pending — physical/network investigation of access point required"
    elif cat == "connectivity" and "is down" in desc_l:
        resolution = "Pending — ISP/network team investigation required"
    else:
        resolution = "Pending — manual investigation required"

    # ── Status ───────────────────────────────────────────────────────────────
    if auto_resolved:
        status = "Resolved"
    elif "reboot required" in desc_l or "suspended" in desc_l:
        status = "Pending"
    else:
        status = "Open"

    # ── Logged By & Follow-up ────────────────────────────────────────────────
    # Populated only when system handled it automatically.
    # Left empty when manual response is needed (human must fill in later).
    if auto_resolved:
        logged_by = "Sophos Central (Auto)"
        follow_up = "Monitor for recurrence"
    else:
        logged_by = ""
        follow_up = ""

    return {
        "incident_type": incident_type,
        "source_of_issue": source_of_issue,
        "affected_system": affected,
        "initial_response": initial_response,
        "resolution": resolution,
        "status": status,
        "logged_by": logged_by,
        "follow_up": follow_up,
        "auto_resolved": auto_resolved,
    }


def _parse_dt(value) -> datetime | None:
    """Parse an ISO date string ('YYYY-MM-DD' or 'YYYY-MM-DDTHH:MM:SS') to a
    timezone-aware datetime, or return None. asyncpg requires actual datetime
    objects — it cannot cast strings to timestamptz itself."""
    if not value:
        return None
    try:
        # Accept YYYY-MM-DD or full ISO string
        if len(value) == 10:
            value += "T00:00:00"
        dt = datetime.fromisoformat(value.rstrip("Z"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt
    except Exception:
        return None

router = APIRouter(prefix="/integrations/sophos", tags=["Integrations - Sophos"])


@router.get("/dashboard")
async def sophos_dashboard(
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    plugin = await db.execute(text("""
        SELECT status, is_enabled, last_sync_at, last_error
        FROM integration_plugins WHERE plugin_type = 'sophos'
    """))
    p = plugin.fetchone()

    stats = await db.execute(text("""
        SELECT
            COUNT(*) AS total,
            COUNT(*) FILTER (WHERE health_status = 'good') AS healthy,
            COUNT(*) FILTER (WHERE health_status = 'suspicious') AS suspicious,
            COUNT(*) FILTER (WHERE health_status = 'bad') AS unhealthy,
            COUNT(*) FILTER (WHERE tamper_protection = TRUE) AS tamper_on,
            COUNT(*) FILTER (WHERE last_seen IS NULL OR last_seen < NOW() - INTERVAL '7 days') AS stale,
            COUNT(*) FILTER (WHERE os_name ILIKE ANY(ARRAY[
                '%Windows 7%','%Windows XP%','%Windows Vista%',
                '%Server 2003%','%Server 2008%','%Server 2012%','%Windows 8%'
            ])) AS eol
        FROM sophos_endpoints
    """))
    s = stats.fetchone()

    alerts = await db.execute(text("""
        SELECT COUNT(*) AS total,
               COUNT(*) FILTER (WHERE severity = 'high') AS high_count,
               COUNT(*) FILTER (WHERE severity = 'medium') AS medium_count
        FROM sophos_alerts
    """))
    a = alerts.fetchone()

    categories = await db.execute(text("""
        SELECT category, COUNT(*) AS cnt
        FROM sophos_alerts
        GROUP BY category ORDER BY cnt DESC
    """))

    # License expiry warnings (expires within 30 days or already expired)
    licenses = await db.execute(text("""
        SELECT
            COUNT(*) FILTER (WHERE expires_at IS NOT NULL AND expires_at < NOW() + INTERVAL '30 days' AND expires_at > NOW()) AS expiring_soon,
            COUNT(*) FILTER (WHERE expires_at IS NOT NULL AND expires_at <= NOW()) AS expired,
            COUNT(*) AS total
        FROM sophos_licenses
    """))
    lic = licenses.fetchone()

    return {
        "plugin": {
            "status": p.status if p else "disconnected",
            "is_enabled": p.is_enabled if p else False,
            "last_sync_at": p.last_sync_at.isoformat() if p and p.last_sync_at else None,
            "last_error": p.last_error if p else None,
        },
        "endpoints": {
            "total": s.total or 0,
            "healthy": s.healthy or 0,
            "suspicious": s.suspicious or 0,
            "unhealthy": s.unhealthy or 0,
            "tamper_protected": s.tamper_on or 0,
            "stale": s.stale or 0,
            "eol": s.eol or 0,
        },
        "alerts": {
            "total": a.total or 0,
            "high": a.high_count or 0,
            "medium": a.medium_count or 0,
            "categories": [{"category": r.category or "unknown", "count": r.cnt}
                           for r in categories.fetchall()],
        },
        "licenses": {
            "total": lic.total or 0,
            "expiring_soon": lic.expiring_soon or 0,
            "expired": lic.expired or 0,
        },
    }


@router.get("/endpoints")
async def list_endpoints(
    health_status: str = None,
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    where = "WHERE health_status = :hs" if health_status else ""
    params = {"hs": health_status} if health_status else {}
    rows = await db.execute(text(f"""
        SELECT endpoint_id, hostname, health_status, os_name, ip_address,
               last_seen, tamper_protection, group_name
        FROM sophos_endpoints
        {where}
        ORDER BY
            CASE health_status WHEN 'bad' THEN 0 WHEN 'suspicious' THEN 1 ELSE 2 END,
            hostname
    """), params)
    return [
        {
            "endpoint_id": r.endpoint_id,
            "hostname": r.hostname,
            "health_status": r.health_status,
            "os_name": r.os_name,
            "ip_address": r.ip_address,
            "last_seen": r.last_seen.isoformat() if r.last_seen else None,
            "tamper_protection": r.tamper_protection,
            "group_name": r.group_name,
        }
        for r in rows.fetchall()
    ]


@router.get("/alerts")
async def list_alerts(
    severity: str = None,
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    where = "WHERE severity = :sev" if severity else ""
    params = {"sev": severity} if severity else {}
    rows = await db.execute(text(f"""
        SELECT alert_id, severity, category, description, endpoint_hostname, raised_at
        FROM sophos_alerts
        {where}
        ORDER BY raised_at DESC NULLS LAST
        LIMIT 200
    """), params)
    return [
        {
            "alert_id": r.alert_id,
            "severity": r.severity,
            "category": r.category,
            "description": r.description,
            "endpoint_hostname": r.endpoint_hostname,
            "raised_at": r.raised_at.isoformat() if r.raised_at else None,
        }
        for r in rows.fetchall()
    ]


@router.get("/licenses")
async def list_licenses(
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    rows = await db.execute(text("""
        SELECT license_id, product_name, license_type,
               starts_at, expires_at, quantity, used_quantity,
               CASE
                   WHEN expires_at IS NULL THEN 'perpetual'
                   WHEN expires_at <= NOW() THEN 'expired'
                   WHEN expires_at <= NOW() + INTERVAL '30 days' THEN 'expiring_soon'
                   ELSE 'active'
               END AS expiry_status,
               CASE WHEN expires_at IS NOT NULL AND expires_at > NOW()
                    THEN EXTRACT(DAY FROM expires_at - NOW())::int
                    ELSE NULL
               END AS days_remaining
        FROM sophos_licenses
        ORDER BY
            CASE
                WHEN expires_at IS NULL THEN 3
                WHEN expires_at <= NOW() THEN 0
                WHEN expires_at <= NOW() + INTERVAL '30 days' THEN 1
                ELSE 2
            END,
            expires_at ASC NULLS LAST
    """))
    return [
        {
            "license_id": r.license_id,
            "product_name": r.product_name,
            "license_type": r.license_type,
            "starts_at": r.starts_at.isoformat() if r.starts_at else None,
            "expires_at": r.expires_at.isoformat() if r.expires_at else None,
            "quantity": r.quantity or 0,
            "used_quantity": r.used_quantity or 0,
            "expiry_status": r.expiry_status,
            "days_remaining": r.days_remaining,
        }
        for r in rows.fetchall()
    ]


@router.post("/licenses")
async def add_license(
    request: Request,
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    """Manually add a license entry (for when API does not expose license info)."""
    body = await request.json()
    product_name = (body.get("product_name") or "").strip()
    if not product_name:
        raise HTTPException(status_code=400, detail="product_name is required")

    await db.execute(text("""
        INSERT INTO sophos_licenses
            (license_id, product_name, license_type, starts_at, expires_at, quantity, used_quantity)
        VALUES (gen_random_uuid()::text, :product_name, :license_type,
                :starts_at, :expires_at, :quantity, :used_quantity)
    """), {
        "product_name": product_name,
        "license_type": (body.get("license_type") or "subscription").lower(),
        "starts_at": _parse_dt(body.get("starts_at")),
        "expires_at": _parse_dt(body.get("expires_at")),
        "quantity": int(body.get("quantity") or 0),
        "used_quantity": int(body.get("used_quantity") or 0),
    })
    await db.commit()
    return {"ok": True}


@router.delete("/licenses/{license_id}")
async def delete_license(
    license_id: str,
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    await db.execute(text("DELETE FROM sophos_licenses WHERE license_id = :lid"), {"lid": license_id})
    await db.commit()
    return {"ok": True}


@router.get("/incident-report")
async def incident_report(
    year: int = Query(default=None, description="Year to report (default: current year)"),
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    """Return structured incident report rows derived from sophos_alerts for the given year."""
    if year is None:
        year = _date.today().year

    # Only include actual security threat categories.
    # Exclude: connectivity (VPN/gateway status), wireless (AP online/offline),
    #          updating (software updates), general (noise).
    THREAT_CATEGORIES = ('malware', 'pua', 'runtimedetections', 'security', 'protection')

    rows = await db.execute(text("""
        SELECT alert_id, severity, category, description, endpoint_hostname, raised_at
        FROM sophos_alerts
        WHERE EXTRACT(YEAR FROM raised_at) = :yr
          AND LOWER(category) = ANY(:cats)
        ORDER BY raised_at ASC
    """), {"yr": year, "cats": list(THREAT_CATEGORIES)})

    out = []
    for i, r in enumerate(rows.fetchall(), 1):
        derived = _derive_incident(r.category or "", r.description or "", r.endpoint_hostname or "")
        raised = r.raised_at
        out.append({
            "no": i,
            "date": raised.strftime("%Y-%m-%d") if raised else "",
            "time": raised.strftime("%H:%M:%S") if raised else "",
            "incident_type": derived["incident_type"],
            "description": r.description or "",
            "severity": (r.severity or "").capitalize(),
            "source": "Sophos Central",
            "source_of_issue": derived["source_of_issue"],
            "affected_system": derived["affected_system"],
            "initial_response": derived["initial_response"],
            "resolution": derived["resolution"],
            "status": derived["status"],
            "logged_by": derived["logged_by"],
            "follow_up": derived["follow_up"],
        })
    return {"year": year, "total": len(out), "incidents": out}


@router.get("/incident-report/xlsx")
async def incident_report_xlsx(
    year: int = Query(default=None, description="Year to export"),
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    """Generate incident report XLSX for the given year."""
    import io
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from fastapi.responses import StreamingResponse

    if year is None:
        year = _date.today().year

    # Fetch data (reuse endpoint logic)
    data = await incident_report(year=year, db=db, _=_)
    incidents = data["incidents"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Incident Report {year}"
    ws.sheet_view.showGridLines = False

    def fill(hex_c):  return PatternFill("solid", fgColor=hex_c)
    def font(bold=False, color="FFFFFF", size=10): return Font(bold=bold, color=color, size=size, name="Calibri")
    def align(h="left", wrap=False): return Alignment(horizontal=h, vertical="center", wrap_text=wrap)
    def border():
        s = Side(style="thin", color="334155")
        return Border(left=s, right=s, top=s, bottom=s)

    # ── Title row ─────────────────────────────────────────────────────────
    ws.merge_cells("A1:N1")
    t = ws["A1"]
    t.value = f"Sophos Central — Incident Report {year}"
    t.font = Font(bold=True, color="FFFFFF", size=13, name="Calibri")
    t.fill = fill("1E3A5F")
    t.alignment = align("left")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:N2")
    s = ws["A2"]
    s.value = f"Generated: {_date.today().strftime('%d %B %Y')}   |   Total incidents: {len(incidents)}   |   Source: Sophos Central"
    s.font = Font(color="94A3B8", size=9, name="Calibri")
    s.fill = fill("0F172A")
    s.alignment = align("left")
    ws.row_dimensions[2].height = 16

    ws.row_dimensions[3].height = 8  # spacer

    # ── Column headers ────────────────────────────────────────────────────
    HEADERS = ["#", "Date", "Time", "Incident Type", "Description", "Severity",
               "Source", "Source of Issue", "Affected System / User",
               "Initial Response", "Resolution", "Status", "Logged By", "Follow-up Actions"]
    COL_WIDTHS = [5, 12, 10, 22, 50, 10, 16, 24, 26, 40, 40, 12, 20, 30]

    for col, (h, w) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        c = ws.cell(row=4, column=col, value=h)
        c.font = Font(bold=True, color="94A3B8", size=9, name="Calibri")
        c.fill = fill("1E293B")
        c.alignment = Alignment(horizontal="center" if col <= 3 else "left", vertical="center", wrap_text=True)
        c.border = border()
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[4].height = 22
    ws.freeze_panes = "A5"

    # ── Data rows ─────────────────────────────────────────────────────────
    SEV_COLORS   = {"High": "7F1D1D", "Medium": "78350F", "Low": "1E3A5F", "": "1E293B"}
    SEV_FG       = {"High": "FCA5A5", "Medium": "FDE68A", "Low": "93C5FD", "": "94A3B8"}
    STATUS_COLORS = {"Resolved": "14532D", "Pending": "78350F", "Open": "7F1D1D"}
    STATUS_FG     = {"Resolved": "86EFAC", "Pending": "FDE68A", "Open": "FCA5A5"}
    ROW_BG = ("1E293B", "0F172A")

    for i, inc in enumerate(incidents):
        row = 5 + i
        bg = ROW_BG[i % 2]
        sev  = inc["severity"]
        stat = inc["status"]

        values = [
            inc["no"], inc["date"], inc["time"], inc["incident_type"],
            inc["description"], inc["severity"], inc["source"],
            inc["source_of_issue"], inc["affected_system"],
            inc["initial_response"], inc["resolution"], inc["status"],
            inc["logged_by"], inc["follow_up"],
        ]
        for col, val in enumerate(values, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.border = border()

            # Special colouring for severity (col 6) and status (col 12)
            if col == 6 and sev in SEV_COLORS:
                c.fill = fill(SEV_COLORS[sev])
                c.font = Font(bold=True, color=SEV_FG.get(sev, "FFFFFF"), size=9, name="Calibri")
                c.alignment = align("center")
            elif col == 12 and stat in STATUS_COLORS:
                c.fill = fill(STATUS_COLORS[stat])
                c.font = Font(bold=True, color=STATUS_FG.get(stat, "FFFFFF"), size=9, name="Calibri")
                c.alignment = align("center")
            elif col <= 3:
                c.fill = fill(bg)
                c.font = Font(color="94A3B8", size=9, name="Calibri")
                c.alignment = align("center")
            else:
                c.fill = fill(bg)
                c.font = Font(color="E2E8F0" if col in (4, 5, 9) else "CBD5E1", size=9, name="Calibri")
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=(col in (5, 10, 11, 14)))

        ws.row_dimensions[row].height = 42 if inc["description"] and len(inc["description"]) > 80 else 28

    # ── Stream response ───────────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"sophos_incident_report_{year}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.put("/licenses/{license_id}")
async def update_license(
    license_id: str,
    request: Request,
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    body = await request.json()
    await db.execute(text("""
        UPDATE sophos_licenses SET
            product_name  = :product_name,
            license_type  = :license_type,
            starts_at     = :starts_at,
            expires_at    = :expires_at,
            quantity      = :quantity,
            used_quantity = :used_quantity,
            synced_at     = NOW()
        WHERE license_id = :lid
    """), {
        "lid": license_id,
        "product_name": (body.get("product_name") or "").strip(),
        "license_type": (body.get("license_type") or "subscription").lower(),
        "starts_at": _parse_dt(body.get("starts_at")),
        "expires_at": _parse_dt(body.get("expires_at")),
        "quantity": int(body.get("quantity") or 0),
        "used_quantity": int(body.get("used_quantity") or 0),
    })
    await db.commit()
    return {"ok": True}
